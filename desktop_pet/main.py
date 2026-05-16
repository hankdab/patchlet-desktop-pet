from __future__ import annotations

import argparse
import csv
import json
import os
import platform
import random
import re
import subprocess
import sys
import threading
import time
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from tempfile import gettempdir
from tkinter import filedialog, messagebox
import tkinter as tk

from PIL import Image, ImageTk

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD

    DND_AVAILABLE = True
except Exception:
    TkinterDnD = None
    DND_FILES = None
    DND_AVAILABLE = False


IS_MACOS = sys.platform == "darwin"
IS_WINDOWS = os.name == "nt"
IS_FROZEN = bool(getattr(sys, "frozen", False))


def app_resource_dir() -> Path:
    if IS_FROZEN and IS_MACOS:
        return Path(sys.executable).resolve().parents[1] / "Resources"
    if IS_FROZEN:
        return Path(getattr(sys, "_MEIPASS", Path(sys.executable).resolve().parent))
    return Path(__file__).resolve().parent


def app_data_dir() -> Path:
    if IS_MACOS:
        return Path.home() / "Library" / "Application Support" / "Patchlet"
    if IS_WINDOWS:
        return Path(os.environ.get("APPDATA", Path.home() / "AppData" / "Roaming")) / "Patchlet"
    return Path(os.environ.get("XDG_DATA_HOME", Path.home() / ".local" / "share")) / "patchlet"


APP_DIR = app_resource_dir()
SOURCE_DIR = Path(__file__).resolve().parent
WORKSPACE = SOURCE_DIR.parent
DEFAULT_SPRITESHEET_CANDIDATES = [
    APP_DIR / "assets" / "spritesheet.webp",
    SOURCE_DIR / "assets" / "spritesheet.webp",
    Path.home() / ".codex" / "pets" / "patchlet" / "spritesheet.webp",
    WORKSPACE / "tmp" / "hatch-pet" / "patchlet" / "final" / "spritesheet.webp",
]
STAGE_SPRITESHEET_CANDIDATES = {
    "base": DEFAULT_SPRITESHEET_CANDIDATES,
    "evolved": [
        APP_DIR / "assets" / "patchlet-evolved.webp",
        SOURCE_DIR / "assets" / "patchlet-evolved.webp",
        WORKSPACE / "tmp" / "hatch-pet" / "patchlet-evolved" / "final" / "spritesheet.webp",
        Path.home() / ".codex" / "pets" / "patchlet-evolved" / "spritesheet.webp",
    ],
    "ultimate": [
        APP_DIR / "assets" / "patchlet-ultimate.webp",
        SOURCE_DIR / "assets" / "patchlet-ultimate.webp",
        WORKSPACE / "tmp" / "hatch-pet" / "patchlet-ultimate" / "final" / "spritesheet.webp",
        Path.home() / ".codex" / "pets" / "patchlet-ultimate" / "spritesheet.webp",
    ],
}
STAGE_LABELS = {
    "base": "小补丁",
    "evolved": "进化小补丁",
    "ultimate": "终极小补丁",
}
REPORTS_DIR = (app_data_dir() if IS_FROZEN else SOURCE_DIR) / "reports"
TRANSPARENT_KEY = "#ff00ff"
CELL_WIDTH = 192
CELL_HEIGHT = 208
ROWS = {
    "idle": (0, 6),
    "running-right": (1, 8),
    "running-left": (2, 8),
    "waving": (3, 4),
    "jumping": (4, 5),
    "failed": (5, 8),
    "waiting": (6, 6),
    "running": (7, 6),
    "review": (8, 6),
}


def open_path(path: Path) -> None:
    if IS_WINDOWS:
        os.startfile(str(path))  # type: ignore[attr-defined]
        return
    if IS_MACOS:
        subprocess.Popen(["open", str(path)])
        return
    subprocess.Popen(["xdg-open", str(path)])


def normalize_drop_paths(data: str, root: tk.Tk) -> list[Path]:
    # TkDND returns a Tcl list. Paths containing spaces are wrapped in braces,
    # and file:// URLs are common on macOS.
    out: list[Path] = []
    for item in root.tk.splitlist(data):
        text = str(item)
        if text.startswith("file://"):
            from urllib.parse import unquote, urlparse

            parsed = urlparse(text)
            text = unquote(parsed.path)
        out.append(Path(text))
    return out


@dataclass
class DocumentSummary:
    path: Path
    kind: str
    size_bytes: int
    modified: str
    title: str
    characters: int
    words: int
    keywords: list[str]
    headings: list[str]
    excerpt: str
    report_path: Path | None = None
    error: str | None = None


def locate_spritesheet() -> Path:
    for candidate in DEFAULT_SPRITESHEET_CANDIDATES:
        if candidate.is_file():
            return candidate
    checked = "\n".join(str(path) for path in DEFAULT_SPRITESHEET_CANDIDATES)
    raise FileNotFoundError(f"Could not find Patchlet spritesheet. Checked:\n{checked}")


def locate_stage_spritesheets() -> dict[str, Path]:
    located: dict[str, Path] = {}
    for stage, candidates in STAGE_SPRITESHEET_CANDIDATES.items():
        for candidate in candidates:
            if candidate.is_file():
                located[stage] = candidate
                break
    if "base" not in located:
        located["base"] = locate_spritesheet()
    return located


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def first_existing_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="replace")


def extract_docx(path: Path) -> tuple[str, list[str]]:
    try:
        import docx

        doc = docx.Document(str(path))
        blocks: list[str] = []
        headings: list[str] = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            blocks.append(text)
            style_name = getattr(paragraph.style, "name", "") or ""
            if style_name.lower().startswith("heading"):
                headings.append(text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))
        return "\n".join(blocks), headings
    except Exception:
        # Fallback keeps DOCX readable even if python-docx is unavailable or the file is odd.
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
        text = re.sub(r"<[^>]+>", " ", xml)
        return clean_text(text), []


def extract_pdf(path: Path) -> tuple[str, list[str]]:
    import pypdf

    reader = pypdf.PdfReader(str(path))
    pages = []
    for page in reader.pages[:30]:
        pages.append(page.extract_text() or "")
    suffix = ""
    if len(reader.pages) > 30:
        suffix = f"\n\n[Only the first 30 of {len(reader.pages)} pages were read.]"
    return "\n\n".join(pages) + suffix, []


def extract_xlsx(path: Path) -> tuple[str, list[str]]:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[str] = []
    headings: list[str] = list(workbook.sheetnames)
    for sheet_name in workbook.sheetnames[:8]:
        sheet = workbook[sheet_name]
        blocks.append(f"[Sheet: {sheet_name}]")
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                blocks.append(" | ".join(values))
            if index >= 80:
                blocks.append("[Sheet truncated after 80 rows]")
                break
    return "\n".join(blocks), headings


def extract_csv_text(path: Path) -> tuple[str, list[str]]:
    raw = first_existing_text(path)
    rows = []
    for index, row in enumerate(csv.reader(raw.splitlines()), start=1):
        if row:
            rows.append(" | ".join(cell.strip() for cell in row))
        if index >= 120:
            rows.append("[CSV truncated after 120 rows]")
            break
    return "\n".join(rows), []


def extract_text(path: Path) -> tuple[str, list[str], str]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".rst", ".log"}:
        return first_existing_text(path), [], suffix.lstrip(".")
    if suffix == ".json":
        raw = first_existing_text(path)
        try:
            return json.dumps(json.loads(raw), ensure_ascii=False, indent=2), [], "json"
        except Exception:
            return raw, [], "json"
    if suffix == ".csv":
        text, headings = extract_csv_text(path)
        return text, headings, "csv"
    if suffix == ".docx":
        text, headings = extract_docx(path)
        return text, headings, "docx"
    if suffix == ".pdf":
        text, headings = extract_pdf(path)
        return text, headings, "pdf"
    if suffix in {".xlsx", ".xlsm"}:
        text, headings = extract_xlsx(path)
        return text, headings, "spreadsheet"
    raise ValueError(f"Unsupported file type: {suffix or 'no extension'}")


def choose_title(path: Path, text: str, headings: list[str]) -> str:
    if headings:
        return headings[0][:120]
    for line in text.splitlines():
        line = line.strip(" #\t-")
        if len(line) >= 2:
            return line[:120]
    return path.stem


def keyword_list(text: str) -> list[str]:
    stopwords = {
        "the",
        "and",
        "for",
        "with",
        "from",
        "that",
        "this",
        "you",
        "are",
        "was",
        "were",
        "have",
        "has",
        "not",
        "but",
        "all",
        "can",
        "will",
        "your",
        "about",
        "into",
    }
    tokens = re.findall(r"[A-Za-z][A-Za-z0-9_-]{2,}|[\u4e00-\u9fff]{2,6}", text)
    counts = Counter(token.lower() for token in tokens if token.lower() not in stopwords)
    return [token for token, _ in counts.most_common(8)]


def summarize_document(path: Path) -> DocumentSummary:
    stat = path.stat()
    modified = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(stat.st_mtime))
    try:
        text, headings, kind = extract_text(path)
        text = clean_text(text)
        words = len(re.findall(r"[A-Za-z0-9_\u4e00-\u9fff]+", text))
        title = choose_title(path, text, headings)
        excerpt = text[:900] if text else "(No extractable text found.)"
        summary = DocumentSummary(
            path=path,
            kind=kind,
            size_bytes=stat.st_size,
            modified=modified,
            title=title,
            characters=len(text),
            words=words,
            keywords=keyword_list(text),
            headings=headings[:10],
            excerpt=excerpt,
        )
    except Exception as exc:
        summary = DocumentSummary(
            path=path,
            kind=path.suffix.lower().lstrip(".") or "file",
            size_bytes=stat.st_size,
            modified=modified,
            title=path.name,
            characters=0,
            words=0,
            keywords=[],
            headings=[],
            excerpt="",
            error=str(exc),
        )
    summary.report_path = write_report(summary)
    return summary


def format_size(size: int) -> str:
    units = ["B", "KB", "MB", "GB"]
    value = float(size)
    for unit in units:
        if value < 1024 or unit == units[-1]:
            return f"{value:.1f} {unit}" if unit != "B" else f"{int(value)} B"
        value /= 1024
    return f"{size} B"


def write_report(summary: DocumentSummary) -> Path:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", summary.path.stem).strip("_") or "document"
    stamp = time.strftime("%Y%m%d-%H%M%S")
    report = REPORTS_DIR / f"{stamp}-{safe_name}.summary.txt"
    lines = [
        f"File: {summary.path}",
        f"Type: {summary.kind}",
        f"Size: {format_size(summary.size_bytes)}",
        f"Modified: {summary.modified}",
        f"Title: {summary.title}",
        f"Characters: {summary.characters}",
        f"Words/tokens: {summary.words}",
        "",
    ]
    if summary.error:
        lines.extend(["Error:", summary.error, ""])
    else:
        if summary.headings:
            lines.extend(["Headings/sheets:", *[f"- {heading}" for heading in summary.headings], ""])
        if summary.keywords:
            lines.extend(["Keywords:", ", ".join(summary.keywords), ""])
        lines.extend(["Extract:", summary.excerpt])
    report.write_text("\n".join(lines), encoding="utf-8")
    return report


def bubble_text(summary: DocumentSummary) -> str:
    if summary.error:
        return f"读不动这个文件：\n{summary.path.name}\n\n{summary.error}"
    keywords = ", ".join(summary.keywords[:5]) if summary.keywords else "暂无"
    excerpt = re.sub(r"\s+", " ", summary.excerpt).strip()
    if len(excerpt) > 220:
        excerpt = excerpt[:220].rstrip() + "..."
    return (
        f"读完啦：{summary.path.name}\n"
        f"{summary.kind} · {format_size(summary.size_bytes)} · {summary.words} 个词/片段\n"
        f"标题：{summary.title}\n"
        f"关键词：{keywords}\n\n"
        f"{excerpt}"
    )


class PatchletApp:
    def __init__(self, spritesheet: Path):
        root_cls = TkinterDnD.Tk if DND_AVAILABLE else tk.Tk
        self.root = root_cls()
        self.root.title("小补丁")
        self.root.configure(bg=TRANSPARENT_KEY)
        self.root.overrideredirect(True)
        self.root.attributes("-topmost", True)
        self.configure_transparency(self.root)

        self.scale = 0.78
        self.stage_paths = locate_stage_spritesheets()
        if spritesheet:
            self.stage_paths["base"] = spritesheet
        self.stage_order = ["base", "evolved", "ultimate"]
        self.stage_index = 0
        self.stage_name = self.stage_order[self.stage_index]
        self.stage_sprites = {
            stage: self.load_sprites(path)
            for stage, path in self.stage_paths.items()
        }
        self.sprites = self.stage_sprites[self.stage_name]
        self.state = "idle"
        self.frame_index = 0
        self.tick = 0
        self.paused = False
        self.dragging = False
        self.busy = False
        self.last_report: Path | None = None
        self.vx = random.choice([-3, 3])
        self.vy = 0
        self.x = 120
        self.y = 120
        self.edge = "bottom"
        self.edge_loops = 0
        self.edge_step = 4
        self.state_until: float | None = None
        self.current_alpha: float | None = None
        self.running_alpha = 0.68
        self.opaque_alpha = 1.0
        self.hover_margin = 90

        first_image = self.sprites["idle"][0]
        self.label = tk.Label(self.root, image=first_image, bg=TRANSPARENT_KEY, bd=0, highlightthickness=0)
        self.label.pack()
        self.root.geometry(f"{first_image.width()}x{first_image.height()}+{self.x}+{self.y}")

        self.bubble = tk.Toplevel(self.root)
        self.bubble.withdraw()
        self.bubble.overrideredirect(True)
        self.bubble.attributes("-topmost", True)
        self.bubble.configure(bg="#202124")
        self.bubble_label = tk.Label(
            self.bubble,
            text="",
            justify="left",
            wraplength=420,
            padx=12,
            pady=10,
            bg="#202124",
            fg="#f8f9fa",
            font=("Microsoft YaHei UI", 9),
        )
        self.bubble_label.pack()

        self.menu = tk.Menu(self.root, tearoff=0)
        self.menu.add_command(label="暂停/继续闲逛", command=self.toggle_pause)
        self.menu.add_command(label="选择文件给小补丁读", command=self.pick_file)
        self.menu.add_command(label="打开上一份报告", command=self.open_last_report)
        self.menu.add_separator()
        self.menu.add_command(label="退出小补丁", command=self.root.destroy)

        self.bind_events()
        self.place_initially()
        if DND_AVAILABLE:
            self.show_bubble("把文档扔到我身上，我来读。", ms=4200)
        else:
            self.show_bubble("拖拽组件没装好时，可以右键选文件给我读。", ms=5200)
        self.root.after(90, self.animate)
        self.root.after(90, self.wander)

    def configure_transparency(self, window: tk.Tk | tk.Toplevel) -> None:
        if IS_WINDOWS:
            try:
                window.wm_attributes("-transparentcolor", TRANSPARENT_KEY)
                return
            except tk.TclError:
                pass
        if IS_MACOS:
            try:
                window.wm_attributes("-transparent", True)
            except tk.TclError:
                pass
        window.attributes("-alpha", 0.96)

    def load_sprites(self, spritesheet: Path) -> dict[str, list[ImageTk.PhotoImage]]:
        atlas = Image.open(spritesheet).convert("RGBA")
        sprites: dict[str, list[ImageTk.PhotoImage]] = {}
        for state, (row, count) in ROWS.items():
            frames = []
            for column in range(count):
                box = (
                    column * CELL_WIDTH,
                    row * CELL_HEIGHT,
                    (column + 1) * CELL_WIDTH,
                    (row + 1) * CELL_HEIGHT,
                )
                frame = atlas.crop(box)
                width = max(1, round(CELL_WIDTH * self.scale))
                height = max(1, round(CELL_HEIGHT * self.scale))
                frame = frame.resize((width, height), Image.Resampling.NEAREST)
                frames.append(ImageTk.PhotoImage(frame))
            sprites[state] = frames
        return sprites

    def bind_events(self) -> None:
        for widget in (self.root, self.label):
            widget.bind("<ButtonPress-1>", self.start_drag)
            widget.bind("<B1-Motion>", self.drag)
            widget.bind("<ButtonRelease-1>", self.end_drag)
            widget.bind("<Button-3>", self.show_menu)
            widget.bind("<Button-2>", self.show_menu)
            widget.bind("<Control-Button-1>", self.show_menu)
            widget.bind("<Double-Button-1>", lambda _event: self.set_state("waving", 1200))
        if DND_AVAILABLE:
            for widget in (self.root, self.label):
                widget.drop_target_register(DND_FILES)
                widget.dnd_bind("<<Drop>>", self.on_drop)

    def place_initially(self) -> None:
        self.root.update_idletasks()
        width = self.sprites["idle"][0].width()
        height = self.sprites["idle"][0].height()
        screen_w = self.root.winfo_screenwidth()
        screen_h = self.root.winfo_screenheight()
        self.x = max(0, screen_w - width - 80)
        self.y = max(0, screen_h - height - 90)
        self.edge = "bottom"
        self.root.geometry(f"{width}x{height}+{self.x}+{self.y}")

    def show_menu(self, event: tk.Event) -> None:
        self.menu.tk_popup(event.x_root, event.y_root)

    def toggle_pause(self) -> None:
        self.paused = not self.paused
        if self.paused:
            self.set_state("idle")
        self.show_bubble("我先蹲一会儿。" if self.paused else "继续巡逻。", ms=2200)

    def pick_file(self) -> None:
        file_path = filedialog.askopenfilename(
            title="选择一个文件给小补丁读",
            filetypes=[
                ("Documents", "*.txt *.md *.json *.csv *.docx *.pdf *.xlsx *.xlsm"),
                ("All files", "*.*"),
            ],
        )
        if file_path:
            self.process_file(Path(file_path))

    def open_last_report(self) -> None:
        if self.last_report and self.last_report.is_file():
            open_path(self.last_report)
        else:
            messagebox.showinfo("小补丁", "还没有读过文件。")

    def start_drag(self, event: tk.Event) -> None:
        self.dragging = True
        self.set_state("idle")
        self.drag_offset_x = event.x
        self.drag_offset_y = event.y

    def drag(self, event: tk.Event) -> None:
        self.x = event.x_root - self.drag_offset_x
        self.y = event.y_root - self.drag_offset_y
        self.root.geometry(f"+{self.x}+{self.y}")
        self.move_bubble()

    def end_drag(self, _event: tk.Event) -> None:
        self.dragging = False

    def on_drop(self, event: tk.Event) -> None:
        paths = normalize_drop_paths(event.data, self.root)
        files = [path for path in paths if path.is_file()]
        if not files:
            self.show_bubble("这次没有接到文件。", ms=2400)
            return
        self.process_file(files[0])

    def process_file(self, path: Path) -> None:
        if self.busy:
            self.show_bubble("我还在读上一份，稍等一下。", ms=2400)
            return
        self.busy = True
        self.paused = True
        self.set_state("review")
        self.show_bubble(f"正在读：{path.name}")

        def worker() -> None:
            summary = summarize_document(path)
            self.root.after(0, lambda: self.finish_reading(summary))

        threading.Thread(target=worker, daemon=True).start()

    def finish_reading(self, summary: DocumentSummary) -> None:
        self.busy = False
        self.paused = False
        self.last_report = summary.report_path
        self.set_state("failed" if summary.error else "waving", 1800)
        self.show_bubble(bubble_text(summary), ms=15000)

    def set_state(self, state: str, duration_ms: int | None = None) -> None:
        if state not in self.sprites:
            return
        self.state = state
        self.frame_index = 0
        self.state_until = time.monotonic() + duration_ms / 1000 if duration_ms else None

    def reload_stage_if_available(self, stage: str) -> bool:
        if stage in self.stage_sprites:
            return True
        paths = locate_stage_spritesheets()
        path = paths.get(stage)
        if not path:
            return False
        self.stage_paths[stage] = path
        self.stage_sprites[stage] = self.load_sprites(path)
        return True

    def evolve_if_ready(self) -> None:
        target_index = min(self.edge_loops, len(self.stage_order) - 1)
        if target_index <= self.stage_index:
            return
        target_stage = self.stage_order[target_index]
        if not self.reload_stage_if_available(target_stage):
            self.show_bubble(f"{STAGE_LABELS[target_stage]}还在孵化，先继续跑。", ms=3200)
            return
        self.stage_index = target_index
        self.stage_name = target_stage
        self.sprites = self.stage_sprites[target_stage]
        self.set_state("jumping", 1800)
        self.show_bubble(f"{STAGE_LABELS[target_stage]}进化完成。", ms=4200)

    def move_along_desktop_edge(self, screen_w: int, screen_h: int, width: int, height: int) -> None:
        left = 8
        top = 8
        right = max(left, screen_w - width - 8)
        bottom = max(top, screen_h - height - 48)

        if self.edge == "bottom":
            self.state = "running-left"
            self.x -= self.edge_step
            self.y = bottom
            if self.x <= left:
                self.x = left
                self.edge = "left"
        elif self.edge == "left":
            self.state = "running"
            self.y -= self.edge_step
            self.x = left
            if self.y <= top:
                self.y = top
                self.edge = "top"
        elif self.edge == "top":
            self.state = "running-right"
            self.x += self.edge_step
            self.y = top
            if self.x >= right:
                self.x = right
                self.edge = "right"
        else:
            self.state = "running"
            self.y += self.edge_step
            self.x = right
            if self.y >= bottom:
                self.y = bottom
                self.edge = "bottom"
                self.edge_loops += 1
                self.evolve_if_ready()

    def animate(self) -> None:
        frames = self.sprites[self.state]
        self.label.configure(image=frames[self.frame_index % len(frames)])
        self.frame_index += 1
        if getattr(self, "state_until", None) and time.monotonic() >= self.state_until:
            self.state_until = None
            self.state = "idle"
            self.frame_index = 0
        self.root.after(140 if self.state.startswith("running") else 180, self.animate)

    def cursor_near_pet(self) -> bool:
        pointer_x = self.root.winfo_pointerx()
        pointer_y = self.root.winfo_pointery()
        width = self.sprites["idle"][0].width()
        height = self.sprites["idle"][0].height()
        return (
            self.x - self.hover_margin <= pointer_x <= self.x + width + self.hover_margin
            and self.y - self.hover_margin <= pointer_y <= self.y + height + self.hover_margin
        )

    def update_alpha(self) -> None:
        running = self.state in {"running-right", "running-left", "running"}
        desired = self.opaque_alpha if self.cursor_near_pet() or not running else self.running_alpha
        if self.current_alpha != desired:
            self.root.attributes("-alpha", desired)
            self.current_alpha = desired

    def wander(self) -> None:
        width = self.sprites["idle"][0].width()
        height = self.sprites["idle"][0].height()
        screen_w = self.root.winfo_screenwidth()
        screen_h = self.root.winfo_screenheight()
        timed_state_active = bool(self.state_until and time.monotonic() < self.state_until)
        can_move = not self.paused and not self.dragging and not self.busy and not timed_state_active
        if can_move:
            if random.random() < 0.012:
                self.set_state(random.choice(["idle", "waiting"]), random.randint(900, 2200))
                self.update_alpha()
                self.root.after(80, self.wander)
                return
            previous_state = self.state
            self.move_along_desktop_edge(screen_w, screen_h, width, height)
            if self.state != previous_state:
                self.frame_index = 0
            self.root.geometry(f"+{int(self.x)}+{int(self.y)}")
            self.move_bubble()
        self.update_alpha()
        self.root.after(80, self.wander)

    def show_bubble(self, text: str, ms: int | None = None) -> None:
        self.bubble_label.configure(text=text)
        self.bubble.deiconify()
        self.move_bubble()
        if ms:
            token = object()
            self.bubble_token = token

            def hide_if_current() -> None:
                if getattr(self, "bubble_token", None) is token:
                    self.bubble.withdraw()

            self.root.after(ms, hide_if_current)

    def move_bubble(self) -> None:
        if not self.bubble.winfo_viewable():
            return
        self.root.update_idletasks()
        self.bubble.update_idletasks()
        bubble_w = self.bubble.winfo_width()
        bubble_h = self.bubble.winfo_height()
        screen_w = self.root.winfo_screenwidth()
        x = max(8, min(screen_w - bubble_w - 8, int(self.x - bubble_w / 2 + 70)))
        y = max(8, int(self.y - bubble_h - 8))
        self.bubble.geometry(f"+{x}+{y}")

    def run(self) -> None:
        self.root.mainloop()


def run_self_test() -> None:
    sheet = locate_spritesheet()
    with Image.open(sheet) as image:
        assert image.size == (1536, 1872), image.size
    sample = SOURCE_DIR / "README.md"
    if not sample.is_file():
        sample = Path(gettempdir()) / "patchlet-self-test.md"
        sample.write_text(
            "# 小补丁桌面宠物\n\n"
            "这是一个用于验证打包版本文档摘要功能的临时文件。"
            "它会确认 spritesheet 能被读取，报告目录可以写入，"
            "并且 Markdown 文本能够被摘要流程正常处理。"
            "如果这段文字能生成报告，说明 macOS app bundle 的基本资源路径可用。",
            encoding="utf-8",
        )
    summary = summarize_document(sample)
    assert summary.characters > 100
    print(json.dumps(
        {
            "ok": True,
            "spritesheet": str(sheet),
            "dnd_available": DND_AVAILABLE,
            "platform": platform.platform(),
            "sample_report": str(summary.report_path),
            "sample_title": summary.title,
        },
        ensure_ascii=False,
        indent=2,
    ))


def inspect_file(path: Path) -> None:
    summary = summarize_document(path)
    print(json.dumps(summary.__dict__ | {"path": str(summary.path), "report_path": str(summary.report_path)}, ensure_ascii=False, indent=2, default=str))


def main() -> None:
    parser = argparse.ArgumentParser(description="Patchlet desktop pet.")
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--inspect", type=Path, help="Read one file and print a JSON summary without opening the pet.")
    parser.add_argument("--spritesheet", type=Path, help="Override spritesheet path.")
    args = parser.parse_args()

    if args.self_test:
        run_self_test()
        return
    if args.inspect:
        inspect_file(args.inspect)
        return

    try:
        spritesheet = args.spritesheet or locate_spritesheet()
        PatchletApp(spritesheet).run()
    except Exception as exc:
        messagebox.showerror("小补丁启动失败", str(exc))
        raise


if __name__ == "__main__":
    main()
